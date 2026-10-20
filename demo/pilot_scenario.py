"""Commercial pilot scenario: a fictitious supplier catalog run through the whole MLA pipeline, measured.

Company "Aurora Autopeças" and every product, code, GTIN, vehicle and compatibility are INVENTED. Vehicle names
are deliberately fake ("Marca Alfa Modelo X1") so no real compatibility is implied.

Writes only under demo/pilot/ (git-ignored): fixtures/, data/ (isolated MLA data folder), output/ and
results.json. Never touches the user's data folder.

Usage:
    .venv/Scripts/python demo/pilot_scenario.py            # generate data + run the pipeline
    .venv/Scripts/python demo/pilot_scenario.py --reuse    # reuse the generated fixtures
"""
from __future__ import annotations

import csv
import ctypes
import json
import os
import platform
import random
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PILOT = ROOT / "demo" / "pilot"
FIXTURES = PILOT / "fixtures"
DATA = PILOT / "data"
OUTPUT = PILOT / "output"
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "demo"))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from make_fixtures import gtin13, make_pdf, product_image  # noqa: E402

SEED = 20261019
CATEGORIES = {
    "FRE": ("Freios", ["Pastilha de freio", "Disco de freio", "Sapata de freio", "Cilindro de roda"]),
    "FIL": ("Filtros", ["Filtro de óleo", "Filtro de ar", "Filtro de combustível", "Filtro de cabine"]),
    "SUS": ("Suspensão", ["Amortecedor", "Bieleta", "Bandeja", "Coxim do amortecedor"]),
    "ILU": ("Iluminação", ["Lâmpada farol", "Lanterna traseira", "Farol de neblina", "Lâmpada de painel"]),
    "ELE": ("Elétrica", ["Bobina de ignição", "Cabo de vela", "Relé auxiliar", "Sensor de rotação"]),
    "MOT": ("Motor", ["Correia dentada", "Tensor da correia", "Junta do cabeçote", "Bomba de óleo"]),
    "ARR": ("Arrefecimento", ["Bomba d'água", "Válvula termostática", "Mangueira do radiador", "Reservatório"]),
    "ACE": ("Acessórios", ["Tapete de borracha", "Capa de volante", "Calha de chuva", "Engate de reboque"]),
}
BRANDS = ["Aurora", "Vetor", "Lumina", "Ferrox"]  # invented
VEHICLES = [(f"Marca {b}", f"Modelo {m}") for b in ("Alfa", "Beta", "Gama") for m in ("X1", "X2", "Y3", "Z4")]
VARIANTS = [("Preto", "PT"), ("Cinza", "CZ"), ("Prata", "PR")]
HEADERS = ["Cód. Produto", "SKU Pai", "Variação", "Descrição do item", "Título para anúncio", "Descrição completa",
           "Categoria", "Subcategoria", "Marca", "Modelo", "EAN", "Preço Venda", "Custo", "Qtd Estoque", "Peso (kg)",
           "Comp (cm)", "Larg (cm)", "Alt (cm)", "Fornecedor", "Cor da variação"]


# --------------------------------------------------------------------------- fixtures
def build_fixtures() -> dict:
    """Write the supplier files and return the ground truth used by the integrity checks."""
    rng = random.Random(SEED)
    if FIXTURES.exists():
        shutil.rmtree(FIXTURES)
    photos = FIXTURES / "fotos"
    photos.mkdir(parents=True)
    truth = {"valid": {}, "errors": {}, "duplicates": [], "incomplete": [], "variants": {}, "no_image": [],
             "images_written": 0, "orphan_images": 0, "corrupted_images": 0, "variant_named_images": 0}
    rows: list[list] = []
    n = 0
    codes = list(CATEGORIES)
    for index in range(300):
        code = codes[index % len(codes)]
        category, names = CATEGORIES[code]
        n += 1
        has_variants = index % 300 >= 260  # the last 40 products have 2-3 variations
        sku = f"AUR-VAR-{index - 259:03d}" if has_variants else f"AUR-{code}-{n:04d}"
        name = f"{names[index % len(names)]} {BRANDS[index % 4]} {100 + index}"
        brand, model = BRANDS[index % 4], f"{code}{100 + index}"
        price = rng.randint(1990, 89990)
        stock = rng.randint(0, 60) if not has_variants else 0
        gtin = gtin13(f"200{index:09d}")  # prefix 200: in-store numbers, never a real product
        weight = round(rng.uniform(0.1, 6.5), 2)
        dims = [rng.randint(5, 60), rng.randint(5, 40), rng.randint(3, 30)]
        title = f"{names[index % len(names)]} {brand} {model} - linha {category}"
        description = (f"{name}. Produto fictício do catálogo de demonstração Aurora Autopeças. Peça nova, embalagem "
                       f"individual, referência {model}. Conferir aplicação antes da compra.")
        if index % 25 == 7:  # incomplete: no title / no description
            title, description = "", ""
            truth["incomplete"].append(sku)
        row = [sku, "", "", name, title, description, category, names[index % len(names)], brand, model, gtin,
               _money(price), _money(int(price * 0.6)), str(stock), _dec(weight), *map(str, dims),
               "Distribuidora Aurora (fictícia)", ""]
        error = None
        if index in (11, 53, 97, 151, 203):  # invalid money text
            row[11] = ["12,3x", "abc", "R$ --", "9,99 reais", "1O,00"][[11, 53, 97, 151, 203].index(index)]
            error = "preço inválido"
        elif index in (17, 71, 133, 187, 229, 241, 31, 89):  # GTIN with a wrong check digit
            gtin = gtin[:-1] + str((int(gtin[-1]) + 1) % 10)
            row[10] = gtin  # imported as received (with a warning) but must block the listing
            truth.setdefault("gtin_invalid", []).append(sku)
        elif index in (23, 119, 173, 211):  # GTIN mangled by a spreadsheet (scientific notation)
            row[10] = "2,00000E+12"
            error = "GTIN em notação científica"
        elif index in (41, 167, 223):  # negative stock
            row[13] = "-3"
            error = "estoque negativo"
        if error:
            truth["errors"][sku] = error
        else:
            truth["valid"][sku] = {"price": price, "stock": stock, "gtin": gtin, "name": name,
                                   "variants": has_variants}
        rows.append(row)
        if has_variants:
            chosen = VARIANTS[: 2 + index % 2]
            truth["variants"][sku] = []
            for vname, vcode in chosen:
                vsku = f"{sku}-{vcode}"
                vstock = rng.randint(0, 25)
                rows.append([vsku, sku, vname, "", "", "", "", "", "", "", "", "", "", str(vstock), "", "", "", "",
                             "", vname])
                truth["variants"][sku].append({"sku": vsku, "stock": vstock})
        # images: most products 1-2 photos named by SKU; some none; a few variant-named ones
        if index % 30 == 13:
            truth["no_image"].append(sku)
            continue
        hue = (index * 37) % 360
        product_image(photos / f"{sku}.jpg", sku, name[:28], hue)
        truth["images_written"] += 1
        if index % 2 == 0:
            product_image(photos / f"{sku}_2.jpg", sku, "vista lateral", hue, 1)
            truth["images_written"] += 1
        if index % 5 == 0:
            product_image(photos / f"{sku}-3.jpg", sku, "embalagem", hue, 2)
            truth["images_written"] += 1
        if has_variants and index % 4 == 0:
            vname, vcode = VARIANTS[0]
            product_image(photos / f"{sku}-{vcode}.jpg", f"{sku}-{vcode}", vname, hue, 2)
            truth["images_written"] += 1
            truth["variant_named_images"] += 1
    # duplicates: same SKU repeated later in the file with another name
    for dup_index in (5, 66, 120, 180, 250, 270):
        original = rows_by_sku(rows)[_sku_for(dup_index)]
        copy = list(original)
        copy[3] = copy[3] + " (linha repetida)"
        rows.append(copy)
        truth["duplicates"].append(copy[0])
    # two rows without SKU
    rows.append(["", "", "", "Item sem código", "Item sem código", "", "Acessórios", "", "Aurora", "", "", "19,90",
                 "", "1", "", "", "", "", "", ""])
    rows.append(["", "", "", "Outro item sem código", "", "", "Motor", "", "Vetor", "", "", "29,90", "", "2", "", "",
                 "", "", "", ""])
    truth["errors"]["(sem SKU) x2"] = "sem SKU"
    # orphan and corrupted images
    for k in range(6):
        product_image(photos / f"AUR-XXX-{9000 + k}.jpg", f"AUR-XXX-{9000 + k}", "sem produto", 200)
        truth["orphan_images"] += 1
    for k in range(3):
        (photos / f"AUR-FIL-{2 + 8 * k:04d}_9.jpg").write_bytes(b"isto nao e uma imagem JPEG")
        truth["corrupted_images"] += 1

    from openpyxl import Workbook
    book = Workbook()
    book.active.title = "Capa"
    book.active["A1"] = "Catálogo Aurora Autopeças — dados fictícios para demonstração"
    sheet = book.create_sheet("Catálogo")
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    book.save(FIXTURES / "catalogo-aurora.xlsx")
    truth["sheet_rows"] = len(rows)

    # supplier PDF with technical data for 28 existing products (2 rows with an empty cell -> human review)
    valid_simple = [s for s, v in truth["valid"].items() if not v["variants"]]
    pdf_skus = valid_simple[:28]
    lines = [(50, 800, "Ficha técnica do fornecedor (fictícia) — Aurora Autopeças")]
    columns = (50, 180, 330, 460)
    y = 760
    lines += [(x, y, t) for x, t in zip(columns, ("Código", "Material", "Garantia", "Cor"))]
    for k, sku in enumerate(pdf_skus):
        y -= 18
        cells = (sku, ["Aço carbono", "Borracha", "Polipropileno", "Alumínio"][k % 4], f"{3 + k % 9} meses",
                 "" if k in (4, 19) else ["Preto", "Cinza", "Prata"][k % 3])
        lines += [(x, y, t) for x, t in zip(columns, cells) if t]
    make_pdf(FIXTURES / "ficha-tecnica-fornecedor.pdf", lines)
    truth["pdf_rows"] = len(pdf_skus)
    truth["pdf_review_rows"] = 2

    # compatibility CSV (fictitious vehicles); 4 unknown SKUs, 2 with inverted years
    compat_rows = [["SKU", "Marca do veículo", "Modelo do veículo", "Ano inicial", "Ano final"]]
    compat_skus = valid_simple[30:90]
    for k, sku in enumerate(compat_skus):
        for j in range(2 + k % 2):
            brand, model = VEHICLES[(k + j) % len(VEHICLES)]
            start = 2008 + (k + j) % 10
            compat_rows.append([sku, brand, model, str(start), str(start + 4)])
    for k in range(4):
        compat_rows.append([f"AUR-NAO-{k:04d}", "Marca Alfa", "Modelo X1", "2015", "2018"])
    compat_rows.append([compat_skus[0], "Marca Beta", "Modelo Y3", "2020", "2012"])
    compat_rows.append([compat_skus[1], "Marca Gama", "Modelo Z4", "2021", "2011"])
    with (FIXTURES / "aplicacoes-veiculos.csv").open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh, delimiter=";").writerows(compat_rows)
    truth["compat_rows"] = len(compat_rows) - 1
    truth["compat_expected_errors"] = 6
    (FIXTURES / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=1), encoding="utf-8")
    return truth


def _sku_for(index: int) -> str:
    code = list(CATEGORIES)[index % len(CATEGORIES)]
    return f"AUR-VAR-{index - 259:03d}" if index >= 260 else f"AUR-{code}-{index + 1:04d}"


def rows_by_sku(rows: list[list]) -> dict[str, list]:
    return {r[0]: r for r in rows if r[0]}


def _money(cents: int) -> str:
    return f"{cents // 100},{cents % 100:02d}"


def _dec(value: float) -> str:
    return f"{value:.2f}".replace(".", ",")


# --------------------------------------------------------------------------- measurement helpers
class _Mem(ctypes.Structure):
    _fields_ = [("cb", ctypes.c_ulong), ("PageFaultCount", ctypes.c_ulong), ("PeakWorkingSetSize", ctypes.c_size_t),
                ("WorkingSetSize", ctypes.c_size_t), ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPagedPoolUsage", ctypes.c_size_t), ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
                ("QuotaNonPagedPoolUsage", ctypes.c_size_t), ("PagefileUsage", ctypes.c_size_t),
                ("PeakPagefileUsage", ctypes.c_size_t)]


def peak_memory_mb() -> float | None:
    if sys.platform != "win32":
        return None
    from ctypes import wintypes
    counters = _Mem()
    counters.cb = ctypes.sizeof(_Mem)
    kernel32 = ctypes.windll.kernel32
    kernel32.GetCurrentProcess.restype = wintypes.HANDLE
    kernel32.K32GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(_Mem), wintypes.DWORD]
    if kernel32.K32GetProcessMemoryInfo(kernel32.GetCurrentProcess(), ctypes.byref(counters), counters.cb):
        return round(counters.PeakWorkingSetSize / 1024 / 1024, 1)
    return None


def machine() -> dict:
    info = {"sistema": f"{platform.system()} {platform.release()} ({platform.version()})",
            "processador": platform.processor(), "nucleos_logicos": os.cpu_count(), "python": platform.python_version()}
    if sys.platform == "win32":
        class _Status(ctypes.Structure):
            _fields_ = [("dwLength", ctypes.c_ulong), ("dwMemoryLoad", ctypes.c_ulong),
                        ("ullTotalPhys", ctypes.c_ulonglong), ("ullAvailPhys", ctypes.c_ulonglong),
                        ("ullTotalPageFile", ctypes.c_ulonglong), ("ullAvailPageFile", ctypes.c_ulonglong),
                        ("ullTotalVirtual", ctypes.c_ulonglong), ("ullAvailVirtual", ctypes.c_ulonglong),
                        ("ullAvailExtendedVirtual", ctypes.c_ulonglong)]
        status = _Status()
        status.dwLength = ctypes.sizeof(_Status)
        ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(status))
        info["ram_total_gb"] = round(status.ullTotalPhys / 1024 ** 3, 1)
    return info


class Timer:
    def __init__(self):
        self.steps: list[dict] = []

    def run(self, name: str, fn):
        start = time.perf_counter()
        value = fn()
        self.steps.append({"etapa": name, "segundos": round(time.perf_counter() - start, 2)})
        print(f"  {name}: {self.steps[-1]['segundos']} s", flush=True)
        return value


# --------------------------------------------------------------------------- pipeline
def run_pipeline(truth: dict) -> dict:
    if DATA.exists():
        shutil.rmtree(DATA)
    if OUTPUT.exists():
        shutil.rmtree(OUTPUT)
    OUTPUT.mkdir(parents=True)
    from PySide6.QtGui import QGuiApplication
    QGuiApplication.instance() or QGuiApplication([])

    from mla.config.paths import AppPaths
    from mla.core.enums import CatalogStatus, ImportAction, Marketplace
    from mla.importing.service import ImportOptions
    from mla.services.context import AppContext
    from mla.services.listing_export import ListingExporter

    t = Timer()
    results: dict = {"maquina": machine(), "etapas": t.steps}
    ctx = AppContext.open(AppPaths(DATA))
    try:
        # 1-3: import the catalog spreadsheet (identifiers, variations, validation of every row)
        analysis = t.run("Importação XLSX: análise", lambda: ctx.imports.analyze(FIXTURES / "catalogo-aurora.xlsx"))
        mapping = analysis.mapping()
        auto_mapped = {col: f for col, f in mapping.items() if f}
        manual = {"Cor da variação": "attr"}
        for column, field in manual.items():
            mapping[column] = field
        results["mapeamento"] = {"colunas": len(analysis.headers), "sugeridas_automaticamente": len(auto_mapped),
                                 "a_confirmar": [s.column for s in analysis.suggestions if s.needs_confirmation],
                                 "ajustes_manuais": manual, "sugestoes": auto_mapped}
        batch = t.run("Importação XLSX: plano e validação por linha",
                      lambda: ctx.imports.plan(analysis, mapping, ImportOptions(match_images_by_sku=False)))
        actions = ctx.imports.repo.counts_by_action(batch.id)
        error_rows = ctx.imports.rows(batch.id, actions=[ImportAction.ERROR], limit=10_000)
        results["analise_planilha"] = {"linhas": sum(actions.values()), "por_acao": actions,
                                       "exemplos_de_erro": sorted({r.issues[0]["message"] for r in error_rows})[:12]}
        report = t.run("Importação XLSX: gravação", lambda: ctx.imports.commit(batch.id))
        results["gravacao_planilha"] = report.counts

        # 4: associate the photo folder by SKU
        images = ctx.imports.analyze_image_folder(FIXTURES / "fotos")
        img_batch = t.run("Associação de imagens: análise", lambda: ctx.imports.plan(images, {}, ImportOptions()))
        img_actions = ctx.imports.repo.counts_by_action(img_batch.id)
        img_report = t.run("Associação de imagens: gravação", lambda: ctx.imports.commit(img_batch.id))
        results["imagens"] = {"plano": img_actions, "gravacao": img_report.counts,
                              "mensagem_plano": ctx.imports.get(img_batch.id).message}

        # supplier PDF (technical data); ambiguous rows stay for a human decision
        pdf = t.run("PDF do fornecedor: análise", lambda: ctx.imports.analyze(FIXTURES / "ficha-tecnica-fornecedor.pdf"))
        pdf_batch = t.run("PDF do fornecedor: plano", lambda: ctx.imports.plan(pdf, pdf.mapping()))
        pdf_actions = ctx.imports.repo.counts_by_action(pdf_batch.id)
        pdf_report = t.run("PDF do fornecedor: gravação", lambda: ctx.imports.commit(pdf_batch.id))
        results["pdf"] = {"mapeamento": pdf.mapping(), "plano": pdf_actions, "gravacao": pdf_report.counts}

        # compatibility base
        compat = ctx.imports.analyze(FIXTURES / "aplicacoes-veiculos.csv", "compatibility")
        compat_batch = t.run("Compatibilidades: plano", lambda: ctx.imports.plan(compat, compat.mapping()))
        compat_actions = ctx.imports.repo.counts_by_action(compat_batch.id)
        compat_report = t.run("Compatibilidades: gravação", lambda: ctx.imports.commit(compat_batch.id))
        results["compatibilidades"] = {"plano": compat_actions, "gravacao": compat_report.counts}

        # 5-6: validation state of the catalog
        summaries = ctx.catalog.search()
        by_status: dict[str, int] = {}
        for s in summaries:
            by_status[s.status.label] = by_status.get(s.status.label, 0) + 1
        results["catalogo"] = {"produtos": len(summaries), "por_situacao": by_status}

        # 8: safe bulk correction (descriptive field only) with preview
        filter_ids = [s.id for s in summaries if s.category == "Filtros"]
        preview = t.run("Edição em lote: pré-visualização",
                        lambda: ctx.catalog.bulk_preview(filter_ids, {"warranty": "90 dias (fictício)"}))
        changed = t.run("Edição em lote: gravação",
                        lambda: ctx.catalog.bulk_update(filter_ids, {"warranty": "90 dias (fictício)"}))
        results["edicao_em_lote"] = {"produtos_selecionados": len(filter_ids), "alteracoes_previstas": len(preview),
                                     "produtos_alterados": changed}

        # 9: prepare listings for the (fictitious) Mercado Livre account; no approval, no publication
        account = ctx.accounts.default_for(Marketplace.MERCADO_LIVRE)
        ready_ids = [s.id for s in ctx.catalog.search() if s.status is CatalogStatus.READY]
        prepared = t.run("Preparação de anúncios (Mercado Livre, conta de teste)",
                         lambda: ctx.listings.prepare(ready_ids, Marketplace.MERCADO_LIVRE, account.id))
        counts = ctx.listings.counts(Marketplace.MERCADO_LIVRE, account.id)
        results["anuncios"] = {"produtos_prontos": len(ready_ids), "criados": len(prepared.created),
                               "por_situacao": {s.label: n for s, n in counts.items() if n}}

        # 10: export the intermediate sheet with photos + the catalog pending report
        rows = ctx.listings.rows(Marketplace.MERCADO_LIVRE, account.id, limit=10_000)
        exporter = ListingExporter(ctx)
        sheet = OUTPUT / "anuncios-mercado-livre.xlsx"
        exported = t.run("Exportação da planilha de anúncios + fotos",
                         lambda: exporter.export_listings([r.listing for r in rows], sheet))
        issues = t.run("Relatório de pendências do catálogo",
                       lambda: exporter.export_catalog_issues(OUTPUT / "pendencias-catalogo.xlsx"))
        results["exportacao"] = {"anuncios": exported.listings, "linhas": exported.rows,
                                 "fotos_copiadas": exported.images_copied, "fotos_nao_encontradas":
                                 exported.images_missing, "com_pendencia_bloqueante": exported.with_errors,
                                 "linhas_relatorio_pendencias": issues.rows,
                                 "produtos_com_pendencia_bloqueante": issues.with_errors}
        results["integridade"] = check_integrity(ctx, truth, sheet)
        results["historico_alteracoes"] = ctx.db.scalar("SELECT COUNT(*) FROM product_changes")
    finally:
        ctx.close()
    results["tempo_total_pipeline_s"] = round(sum(s["segundos"] for s in t.steps), 2)
    results["pico_memoria_mb"] = peak_memory_mb()
    return results


def check_integrity(ctx, truth: dict, sheet: Path) -> dict:
    from openpyxl import load_workbook

    from mla.services.listing_export import LISTING_COLUMNS
    checks: dict[str, object] = {}
    catalog_skus = {s.sku for s in ctx.catalog.search()}
    valid = truth["valid"]
    checks["skus_validos_na_origem"] = len(valid)
    checks["skus_validos_no_catalogo"] = len(set(valid) & catalog_skus)
    checks["skus_com_erro_que_entraram"] = sorted(set(truth["errors"]) & catalog_skus)
    checks["skus_inesperados_no_catalogo"] = sorted(catalog_skus - set(valid))
    price_mismatch, stock_mismatch, gtin_mismatch = [], [], []
    for sku, expected in valid.items():
        pid = ctx.products.id_for_sku(sku)
        if pid is None:
            continue
        product = ctx.products.get(pid)
        if product.price_cents != expected["price"]:
            price_mismatch.append(sku)
        if product.gtin != expected["gtin"]:
            gtin_mismatch.append(sku)
        if expected["variants"]:
            got = sorted((v.sku, v.stock) for v in product.variants)
            want = sorted((v["sku"], v["stock"]) for v in truth["variants"][sku])
            if got != want:
                stock_mismatch.append(sku)
        elif product.stock != expected["stock"]:
            stock_mismatch.append(sku)
    checks["precos_divergentes"] = price_mismatch
    checks["estoques_ou_variacoes_divergentes"] = stock_mismatch
    checks["gtin_divergentes"] = gtin_mismatch
    book = load_workbook(sheet, read_only=True)
    rows = list(book["Anúncios"].iter_rows(min_row=2, values_only=True))
    col = {name: i for i, name in enumerate(LISTING_COLUMNS)}
    sheet_mismatch = []
    for row in rows:
        sku, parent = row[col["SKU"]], row[col["SKU do produto pai"]]
        source = valid.get(parent or sku)
        if source is None:
            sheet_mismatch.append(sku)
            continue
        gtin = row[col["GTIN/EAN"]] or ""
        if gtin and gtin != source["gtin"]:
            sheet_mismatch.append(sku)
        if not parent and row[col["Preço (R$)"]] != _money(source["price"]):
            sheet_mismatch.append(sku)
    invalid_gtin = set(truth.get("gtin_invalid", []))
    exported = {row[col["SKU do produto pai"]] or row[col["SKU"]] for row in rows}
    blocked = [row[col["SKU"]] for row in rows if row[col["SKU"]] in invalid_gtin
               and "dígito verificador inválido" in (row[col["Pendências"]] or "")]
    checks["gtin_invalido_importado_com_aviso"] = len(invalid_gtin & catalog_skus)
    checks["gtin_invalido_exportado"] = len(invalid_gtin & exported)
    checks["gtin_invalido_bloqueado_no_anuncio"] = len(blocked)
    checks["linhas_da_planilha"] = len(rows)
    checks["linhas_da_planilha_divergentes_da_origem"] = sheet_mismatch
    checks["ok"] = not (checks["skus_com_erro_que_entraram"] or checks["skus_inesperados_no_catalogo"]
                        or price_mismatch or stock_mismatch or gtin_mismatch or sheet_mismatch
                        or checks["skus_validos_no_catalogo"] != len(valid)
                        or len(blocked) != len(invalid_gtin & exported))
    return checks


def main() -> int:
    from PySide6.QtGui import QGuiApplication
    QGuiApplication.instance() or QGuiApplication([])
    PILOT.mkdir(parents=True, exist_ok=True)
    reuse = "--reuse" in sys.argv and (FIXTURES / "truth.json").exists()
    start = time.perf_counter()
    if reuse:
        truth = json.loads((FIXTURES / "truth.json").read_text(encoding="utf-8"))
        generation = None
    else:
        print("Gerando dados fictícios…", flush=True)
        truth = build_fixtures()
        generation = round(time.perf_counter() - start, 1)
    print("Executando o fluxo completo…", flush=True)
    results = run_pipeline(truth)
    results["geracao_dos_dados_s"] = generation
    results["dados_gerados"] = {k: (len(v) if isinstance(v, (dict, list)) else v) for k, v in truth.items()}
    (PILOT / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: results[k] for k in ("tempo_total_pipeline_s", "pico_memoria_mb", "integridade")},
                     ensure_ascii=False, indent=1))
    return 0 if results["integridade"]["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
