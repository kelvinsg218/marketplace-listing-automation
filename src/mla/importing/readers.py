"""Readers turn a source file into headers + rows of strings.

CSV and XLSX are streamed (rows are read again on each pass, never all kept in
memory). JSON, XML and PDF are parsed whole by nature. Readers never interpret
values: they only extract text; parsing and validation happen in the planner.

PDF: rows of a table whose cells line up with the header are "safe"; anything
misaligned, or a technical sheet read as "Name: value" lines, is flagged for review.
"""
from __future__ import annotations

import codecs
import csv
import datetime as dt
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from mla.importing.mapping import KNOWN_KEYS, normalize_key

EXTENSIONS = {".csv": "csv", ".txt": "csv", ".tsv": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".json": "json",
              ".xml": "xml", ".pdf": "pdf", ".zip": "zip"}
FORMAT_LABELS = {"csv": "CSV", "xlsx": "Planilha Excel (XLSX)", "json": "JSON", "xml": "XML", "pdf": "PDF",
                 "zip": "Arquivo ZIP"}
MAX_STRUCTURED_BYTES = 200 * 1024 * 1024  # JSON / XML / PDF are parsed whole


class SourceError(ValueError):
    pass


@dataclass
class SourceRow:
    number: int                     # sequence used in reports
    values: dict[str, str]
    where: str = ""                 # how the user finds it: "linha 12", "página 2"
    review: str = ""                # non-empty: why a human must check this row


@dataclass
class SourceInfo:
    format: str
    headers: list[str]
    notes: list[str] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)
    sheet: str = ""
    total_rows: int | None = None


class Source:
    """A readable source. `rows()` can be called more than once."""

    def __init__(self, info: SourceInfo):
        self.info = info

    def rows(self) -> Iterator[SourceRow]:  # pragma: no cover - interface
        raise NotImplementedError

    def preview(self, limit: int = 50) -> list[SourceRow]:
        result = []
        for row in self.rows():
            result.append(row)
            if len(result) >= limit:
                break
        return result


# ------------------------------------------------------------------ detection
def detect_format(path: Path) -> str:
    path = Path(path)
    if not path.is_file():
        raise SourceError(f"Arquivo não encontrado: {path}")
    ext = path.suffix.lower()
    if ext == ".xls":
        raise SourceError("Planilhas .xls (Excel 97-2003) não são lidas; salve como .xlsx ou .csv e importe de novo.")
    with path.open("rb") as fh:
        head = fh.read(4096)
    if not head:
        raise SourceError("O arquivo está vazio.")
    sniffed = _sniff(path, head)
    declared = EXTENSIONS.get(ext)
    if declared is None:
        if sniffed is None:
            raise SourceError(f"Formato não reconhecido: {path.name}. Use CSV, XLSX, JSON, XML, PDF ou ZIP.")
        return sniffed
    if declared in ("xlsx", "zip", "pdf") and sniffed != declared:
        raise SourceError(f"{path.name} não parece ser um {FORMAT_LABELS[declared]} válido.")
    if declared in ("json", "xml") and sniffed not in (declared, None):
        raise SourceError(f"{path.name} não parece ser um {FORMAT_LABELS[declared]} válido.")
    return declared


def _sniff(path: Path, head: bytes) -> str | None:
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        import zipfile
        try:
            with zipfile.ZipFile(path) as zf:
                return "xlsx" if "xl/workbook.xml" in zf.namelist() else "zip"
        except zipfile.BadZipFile:
            return None
    text = head.lstrip(codecs.BOM_UTF8).lstrip()
    if text[:1] in (b"{", b"["):
        return "json"
    if text[:1] == b"<":
        return "xml"
    return None


def open_source(path: Path, fmt: str | None = None, sheet: str | None = None) -> Source:
    path = Path(path)
    fmt = fmt or detect_format(path)
    if fmt == "csv":
        return CsvSource(path)
    if fmt == "xlsx":
        return XlsxSource(path, sheet)
    if fmt in ("json", "xml", "pdf") and path.stat().st_size > MAX_STRUCTURED_BYTES:
        raise SourceError(f"Arquivo {FORMAT_LABELS[fmt]} acima de 200 MB; divida em partes menores.")
    if fmt == "json":
        return JsonSource(path)
    if fmt == "xml":
        return XmlSource(path)
    if fmt == "pdf":
        return PdfSource(path)
    raise SourceError(f"{FORMAT_LABELS.get(fmt, fmt)} não é uma tabela; extraia o conteúdo primeiro.")


def _unique_headers(raw: list[Any]) -> list[str]:
    headers, seen = [], Counter()
    for index, value in enumerate(raw, start=1):
        name = str(value).strip() if value is not None else ""
        name = name or f"Coluna {index}"
        seen[name.lower()] += 1
        if seen[name.lower()] > 1:
            name = f"{name} ({seen[name.lower()]})"
        headers.append(name)
    return headers


# ------------------------------------------------------------------ CSV
class CsvSource(Source):
    def __init__(self, path: Path):
        self.path = path
        self.encoding = _detect_encoding(path)
        with path.open("r", encoding=self.encoding, newline="") as fh:
            sample = fh.read(64 * 1024)
        self.delimiter = _detect_delimiter(sample)
        header_row = None
        with path.open("r", encoding=self.encoding, newline="") as fh:
            for record in csv.reader(fh, delimiter=self.delimiter):
                if any(cell.strip() for cell in record):
                    header_row = record
                    break
        if header_row is None:
            raise SourceError("O CSV não tem cabeçalho nem dados.")
        names = {";": "ponto e vírgula", ",": "vírgula", "\t": "tabulação", "|": "barra vertical"}
        super().__init__(SourceInfo("csv", _unique_headers(header_row), notes=[
            f"Codificação: {'UTF-8' if self.encoding.startswith('utf') else 'Windows-1252 (ANSI)'}",
            f"Separador: {names.get(self.delimiter, self.delimiter)}"]))

    def rows(self) -> Iterator[SourceRow]:
        headers = self.info.headers
        with self.path.open("r", encoding=self.encoding, newline="") as fh:
            reader = csv.reader(fh, delimiter=self.delimiter)
            header_seen = False
            number = 0
            for record in reader:
                if not any(cell.strip() for cell in record):
                    continue
                if not header_seen:
                    header_seen = True
                    continue
                number += 1
                values = {h: (record[i] if i < len(record) else "") for i, h in enumerate(headers)}
                review = ""
                extra = [c for c in record[len(headers):] if c.strip()]
                if extra:
                    review = "Linha com mais colunas que o cabeçalho; os dados podem estar deslocados."
                yield SourceRow(number, values, f"linha {reader.line_num}", review)


def _detect_encoding(path: Path) -> str:
    decoder = codecs.getincrementaldecoder("utf-8")()
    try:
        with path.open("rb") as fh:
            for chunk in iter(lambda: fh.read(1 << 20), b""):
                decoder.decode(chunk)
            decoder.decode(b"", final=True)
    except UnicodeDecodeError:
        return "cp1252"
    return "utf-8-sig"


def _detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        first = sample.split("\n", 1)[0]
        return max(";,\t|", key=first.count)


# ------------------------------------------------------------------ XLSX
def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        text = repr(value)
        if "e" in text.lower():
            text = format(value, ".10f").rstrip("0").rstrip(".")
        return text
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


class XlsxSource(Source):
    def __init__(self, path: Path, sheet: str | None = None):
        self.path = path
        wb = self._open()
        try:
            sheets = list(wb.sheetnames)
            if sheet and sheet not in sheets:
                raise SourceError(f"A planilha não tem a aba “{sheet}”.")
            self.sheet = sheet or sheets[0]
            ws = wb[self.sheet]
            header, self._header_row = None, 0
            for index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if any(_cell_text(v).strip() for v in row):
                    header, self._header_row = row, index
                    break
            total = ws.max_row - self._header_row if ws.max_row and header is not None else None
        finally:
            wb.close()
        if header is None:
            raise SourceError(f"A aba “{self.sheet}” está vazia.")
        width = max(i for i, v in enumerate(header, start=1) if _cell_text(v).strip())
        super().__init__(SourceInfo("xlsx", _unique_headers(list(header[:width])), notes=[f"Aba: {self.sheet}"],
                                    sheets=sheets, sheet=self.sheet, total_rows=total))

    def _open(self):
        from openpyxl import load_workbook  # macros are never loaded nor run (keep_vba=False)
        try:
            return load_workbook(self.path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises several types for corrupt files
            raise SourceError(f"Não foi possível abrir a planilha: {exc}") from exc

    def rows(self) -> Iterator[SourceRow]:
        headers = self.info.headers
        wb = self._open()
        try:
            ws = wb[self.sheet]
            number = 0
            for index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if index <= self._header_row:
                    continue
                cells = [_cell_text(v) for v in row]
                if not any(c.strip() for c in cells):
                    continue
                number += 1
                values = {h: (cells[i] if i < len(cells) else "") for i, h in enumerate(headers)}
                extra = [c for c in cells[len(headers):] if c.strip()]
                review = "Linha com valores além das colunas do cabeçalho." if extra else ""
                yield SourceRow(number, values, f"linha {index}", review)
        finally:
            wb.close()


# ------------------------------------------------------------------ JSON
def _flatten(value: Any, prefix: str, out: dict[str, str]) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            _flatten(child, f"{prefix}.{key}" if prefix else str(key), out)
    elif isinstance(value, list):
        if all(not isinstance(v, (dict, list)) for v in value):
            out[prefix] = " | ".join(_cell_text(v) for v in value if v is not None)
        elif all(isinstance(v, dict) for v in value) and all(
                any(k in v for k in ("file", "url", "src", "path")) for v in value):
            out[prefix] = " | ".join(str(v.get("file") or v.get("url") or v.get("src") or v.get("path") or "")
                                     for v in value)
        else:
            out[prefix] = json.dumps(value, ensure_ascii=False)
    else:
        out[prefix] = _cell_text(value)


def _find_records(data: Any, depth: int = 0) -> list[dict] | None:
    if isinstance(data, list) and data and all(isinstance(v, dict) for v in data):
        return data
    if isinstance(data, dict) and depth < 3:
        for key in ("products", "produtos", "items", "itens", "data", "dados", "records", "registros"):
            if key in data:
                found = _find_records(data[key], depth + 1)
                if found is not None:
                    return found
        for value in data.values():
            found = _find_records(value, depth + 1)
            if found is not None:
                return found
    return None


class _RecordsSource(Source):
    """Source over records already parsed in memory."""

    def __init__(self, fmt: str, records: list[dict[str, str]], notes: list[str], where: str = "registro",
                 reviews: list[str] | None = None):
        headers: list[str] = []
        seen: set[str] = set()
        for record in records:
            for key in record:
                if key not in seen:
                    seen.add(key)
                    headers.append(key)
        self._records = records
        self._where = where
        self._reviews = reviews or [""] * len(records)
        super().__init__(SourceInfo(fmt, headers, notes=notes, total_rows=len(records)))

    def rows(self) -> Iterator[SourceRow]:
        for number, (record, review) in enumerate(zip(self._records, self._reviews), start=1):
            yield SourceRow(number, {h: record.get(h, "") for h in self.info.headers}, f"{self._where} {number}",
                            review)


class JsonSource(_RecordsSource):
    def __init__(self, path: Path):
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except UnicodeDecodeError as exc:
            raise SourceError("O JSON precisa estar em UTF-8.") from exc
        except json.JSONDecodeError as exc:
            raise SourceError(f"JSON inválido (linha {exc.lineno}): {exc.msg}") from exc
        notes = []
        if isinstance(data, dict) and data.get("format") == "mla-catalog":
            records, notes = _mla_records(data)
        else:
            found = _find_records(data)
            if found is None:
                raise SourceError("O JSON não contém uma lista de produtos (objetos).")
            records = []
            for item in found:
                flat: dict[str, str] = {}
                _flatten(item, "", flat)
                records.append(flat)
        if not records:
            raise SourceError("O JSON não tem registros.")
        super().__init__("json", records, notes)


def _decimal_text(value: Any, factor: int) -> str:
    """Integer units (cents, mm) -> Brazilian decimal text the planner parses back exactly."""
    if value in (None, ""):
        return ""
    whole, rest = divmod(int(value), factor)
    digits = len(str(factor)) - 1
    return f"{whole},{rest:0{digits}d}" if digits else str(whole)


def _mla_records(data: dict) -> tuple[list[dict[str, str]], list[str]]:
    """MLA's own export: products plus one row per variation (linked by SKU do produto pai)."""
    records: list[dict[str, str]] = []
    dropped_marketplace = False
    for product in data.get("products") or []:
        if not isinstance(product, dict):
            continue
        flat: dict[str, str] = {}
        for key, value in product.items():
            if key in ("variants", "marketplaces", "status", "source"):
                continue
            if key == "images":
                flat["images"] = " | ".join(str(i.get("file") or "") for i in value or [] if isinstance(i, dict))
            elif key in ("specs", "custom_attributes") and isinstance(value, dict):
                for name, text in value.items():
                    flat[f"{'spec' if key == 'specs' else 'atributo'}_{name}"] = _cell_text(text)
            elif key.endswith("_cents"):
                flat[key] = _decimal_text(value, 100)
            elif key.endswith("_mm"):
                flat[key[:-3] + "_cm"] = _decimal_text(value, 10)
            else:
                flat[key] = _cell_text(value)
        records.append(flat)
        dropped_marketplace = dropped_marketplace or bool(product.get("marketplaces"))
        for variant in product.get("variants") or []:
            if isinstance(variant, dict):
                records.append({"sku": _cell_text(variant.get("sku")), "sku_pai": flat.get("sku", ""),
                                "variacao": _cell_text(variant.get("name")),
                                "price_cents": _decimal_text(variant.get("price_cents"), 100),
                                "stock": _cell_text(variant.get("stock"))})
    notes = ["Formato de exportação do MLA (variações viram linhas ligadas ao SKU pai)."]
    if dropped_marketplace:
        notes.append("Dados por marketplace do arquivo não são importados por aqui; prepare os anúncios depois.")
    return records, notes


# ------------------------------------------------------------------ XML
_ENTITY_RE = re.compile(rb"<!(DOCTYPE|ENTITY)", re.IGNORECASE)


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _xml_flatten(element: ET.Element, prefix: str, out: dict[str, str]) -> None:
    for name, value in element.attrib.items():
        out[f"{prefix}@{_local(name)}" if prefix else f"@{_local(name)}"] = value
    children = list(element)
    if not children:
        if prefix:
            text = (element.text or "").strip()
            out[prefix] = f"{out[prefix]} | {text}" if prefix in out and text else (out.get(prefix) or text)
        return
    for child in children:
        key = f"{prefix}.{_local(child.tag)}" if prefix else _local(child.tag)
        _xml_flatten(child, key, out)


class XmlSource(_RecordsSource):
    def __init__(self, path: Path):
        raw = path.read_bytes()
        if _ENTITY_RE.search(raw):
            raise SourceError("XML com DOCTYPE/ENTITY não é aceito por segurança; exporte o arquivo sem DTD.")
        try:
            root = ET.fromstring(raw)
        except ET.ParseError as exc:
            raise SourceError(f"XML inválido: {exc}") from exc
        elements, tag = _record_elements(root)
        records = []
        for element in elements:
            flat: dict[str, str] = {}
            _xml_flatten(element, "", flat)
            records.append(flat)
        if not any(records):
            raise SourceError("Não foram encontrados registros no XML.")
        super().__init__("xml", records, [f"Registros: elementos <{tag}>"])


def _record_elements(root: ET.Element) -> tuple[list[ET.Element], str]:
    """Records = the most repeated *structured* child elements (with sub-elements or attributes),
    searched a few levels deep. A document without repetition is a single record (the root)."""
    best: tuple[int, int, ET.Element, str] | None = None
    frontier = [(root, 0)]
    while frontier:
        element, depth = frontier.pop(0)
        structured = Counter(_local(c.tag) for c in element if len(c) or c.attrib)
        if structured:
            tag, count = structured.most_common(1)[0]
            if best is None or count > best[0]:
                best = (count, depth, element, tag)
        if depth < 3:
            frontier.extend((child, depth + 1) for child in element)
    if best is None:
        return [root], _local(root.tag)
    count, _, element, tag = best
    if count == 1 and any(len(c) == 0 and not c.attrib for c in element):
        return [element], _local(element.tag)  # one record with its own fields (e.g. <produto><sku>…)
    return [c for c in element if _local(c.tag) == tag], tag


# ------------------------------------------------------------------ PDF
_CELL_RE = re.compile(r"\S+(?: \S+)*")
_KV_RE = re.compile(r"^\s*([^:]{2,60}?)\s*:\s*(\S.*?)\s*$")


class PdfSource(_RecordsSource):
    def __init__(self, path: Path):
        lines = _pdf_lines(path)
        table = _pdf_table(lines)
        if table is not None:
            records, reviews, notes = table
        else:
            pairs = [m.groups() for _, line in lines if (m := _KV_RE.match(line))]
            if len(pairs) < 2:
                raise SourceError("Não foi encontrada no PDF uma tabela com cabeçalho reconhecível nem uma ficha "
                                  "técnica no formato “Campo: valor”. Converta para planilha e importe de novo.")
            record: dict[str, str] = {}
            for key, value in pairs:
                key = key.strip()
                record[key] = f"{record[key]} | {value}" if key in record else value
            records, reviews = [record], ["Ficha técnica lida de PDF: confira cada valor antes de importar."]
            notes = [f"Ficha técnica: {len(record)} campo(s) “Campo: valor” encontrados."]
        super().__init__("pdf", records, notes, where="item", reviews=reviews)


def _pdf_lines(path: Path) -> list[tuple[int, str]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise SourceError("PDF protegido por senha.")
        lines = []
        for page_no, page in enumerate(reader.pages, start=1):
            text = page.extract_text(extraction_mode="layout") or ""
            lines += [(page_no, line.rstrip()) for line in text.splitlines() if line.strip()]
    except SourceError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise SourceError(f"Não foi possível ler o PDF: {exc}") from exc
    if not lines:
        raise SourceError("O PDF não tem texto extraível (provavelmente é digitalizado). OCR não é suportado.")
    return lines


def _spans(line: str) -> list[tuple[int, int, str]]:
    """Cells separated by 2+ spaces, with their character positions."""
    return [(m.start(), m.end(), m.group()) for m in _CELL_RE.finditer(line)]


def _pdf_table(lines: list[tuple[int, str]]):
    header_index = None
    for index, (_, line) in enumerate(lines):
        cells = _spans(line)
        known = sum(1 for _, _, text in cells if normalize_key(text) in KNOWN_KEYS)
        if len(cells) >= 2 and known >= 2:
            header_index = index
            break
    if header_index is None:
        return None
    header = _spans(lines[header_index][1])
    names = _unique_headers([text for _, _, text in header])
    header_key = [normalize_key(text) for _, _, text in header]
    records: list[dict[str, str]] = []
    reviews: list[str] = []
    ignored = 0
    for page, line in lines[header_index + 1:]:
        cells = _spans(line)
        if [normalize_key(text) for _, _, text in cells] == header_key:
            continue  # header repeated on a new page
        if len(cells) < 2:
            ignored += 1
            continue
        columns = [_nearest_column(header, start, end) for start, end, _ in cells]
        values = {name: "" for name in names}
        aligned = len(set(columns)) == len(columns) and columns == sorted(columns)
        for column, (_, _, text) in zip(columns, cells):
            name = names[column]
            values[name] = f"{values[name]} {text}".strip()
        review = ""
        if not aligned or len(cells) > len(header):
            review = f"Página {page}: colunas desalinhadas no PDF; confira os valores."
        elif len(cells) < len(header):
            review = f"Página {page}: linha com células vazias; confira se os valores caíram na coluna certa."
        records.append(values)
        reviews.append(review)
    if not records:
        return None
    notes = [f"Tabela encontrada com {len(names)} colunas e {len(records)} linha(s)."]
    flagged = sum(1 for r in reviews if r)
    if flagged:
        notes.append(f"{flagged} linha(s) com extração ambígua exigem revisão.")
    if ignored:
        notes.append(f"{ignored} linha(s) de texto fora da tabela foram ignoradas.")
    return records, reviews, notes


def _nearest_column(header: list[tuple[int, int, str]], start: int, end: int) -> int:
    best, best_score = 0, None
    for index, (h_start, h_end, _) in enumerate(header):
        overlap = min(end, h_end) - max(start, h_start)
        distance = abs((start + end) / 2 - (h_start + h_end) / 2)
        score = (-overlap if overlap > 0 else 0, distance)
        if best_score is None or score < best_score:
            best, best_score = index, score
    return best
