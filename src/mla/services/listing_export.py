"""Spreadsheet exports for the client: prepared listings and catalog pending issues.

The listings sheet is an *intermediate* sheet: one row per variation, columns named after the fields marketplaces
usually ask for, photos copied to a folder next to it. It is not the official bulk-upload template of Mercado Livre
or Shopee (those are generated per category inside the seller account); filling them is a supervised step.
Nothing is invented: empty data stays empty and only confirmed compatibilities are listed.
"""
from __future__ import annotations

import csv
import logging
import shutil
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from mla.core.enums import Verification
from mla.core.models import Listing
from mla.core.units import format_money, format_scaled

log = logging.getLogger(__name__)

NOTICE = ("Planilha intermediária gerada pelo MLA. Não é o modelo oficial do Mercado Livre nem da Shopee: use-a "
          "para conferir os dados e copiá-los para o modelo oficial baixado na conta do vendedor (ou para o "
          "mapeamento supervisionado). As fotos estão na pasta indicada; o modelo oficial do Mercado Livre pede "
          "links de fotos (Gestor de fotos).")

LISTING_COLUMNS = [
    "Situação no MLA", "Pendências", "Marketplace", "Conta", "SKU", "SKU do produto pai", "Variação",
    "Atributos da variação", "Título", "Descrição", "Preço (R$)", "Estoque", "Condição", "GTIN/EAN", "Marca",
    "Modelo", "Categoria (catálogo)", "Subcategoria", "ID da categoria no marketplace", "Tipo de anúncio",
    "Peso (g)", "Comprimento (cm)", "Largura (cm)", "Altura (cm)", "Fotos (arquivos)", "Atributos",
    "Compatibilidades confirmadas", "Compatibilidades não verificadas (não exportadas)", "Origem dos dados",
]
ISSUE_COLUMNS = ["SKU", "Produto", "Situação", "Gravidade", "Campo", "Pendência"]
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


@dataclass
class ExportResult:
    path: Path
    rows: int = 0
    listings: int = 0
    images_copied: int = 0
    images_missing: int = 0
    photos_dir: Path | None = None
    by_status: Counter = field(default_factory=Counter)
    with_errors: int = 0

    def summary(self) -> str:
        text = f"{self.listings} anúncio(s) em {self.rows} linha(s)"
        if self.photos_dir is not None:
            text += f", {self.images_copied} foto(s) copiada(s)"
        if self.images_missing:
            text += f", {self.images_missing} foto(s) não encontrada(s)"
        if self.with_errors:
            text += f"; {self.with_errors} anúncio(s) com pendência bloqueante"
        return text + "."


def _safe(value):
    """CSV: spreadsheet apps run cells that start with = + - @ as formulas; keep client text as plain text."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def _cm(mm: int | None) -> str:
    return format_scaled(mm, 10) if mm else ""


class SpreadsheetWriter:
    """Minimal XLSX/CSV writer: sheets of rows, text never turned into formulas."""

    def __init__(self, path: Path):
        self.path = path
        self.sheets: list[tuple[str, list[str], list[list]]] = []

    def add(self, title: str, columns: list[str], rows: list[list]) -> None:
        self.sheets.append((title, columns, rows))

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.suffix.lower() == ".csv":
            _, columns, rows = self.sheets[0]
            with self.path.open("w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.writer(fh, delimiter=";")
                writer.writerow(columns)
                for row in rows:
                    writer.writerow([_safe(v) if isinstance(v, str) else ("" if v is None else v) for v in row])
            return
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        book = Workbook()
        book.remove(book.active)
        for title, columns, rows in self.sheets:
            sheet = book.create_sheet(title[:31])
            sheet.append(columns)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append(["" if v is None else v for v in row])
            for cells in sheet.iter_rows(min_row=2):
                for cell in cells:
                    if cell.data_type == "f":  # openpyxl turns "=..." into a formula; keep it as text
                        cell.data_type = "s"
            for index, name in enumerate(columns, start=1):
                longest = max([len(str(name))] + [len(str(r[index - 1] or "")) for r in rows[:200]])
                sheet.column_dimensions[get_column_letter(index)].width = min(max(12, longest + 2), 60)
            sheet.freeze_panes = "A2"
        tmp = self.path.with_suffix(self.path.suffix + ".tmp")
        book.save(tmp)
        tmp.replace(self.path)


class ListingExporter:
    def __init__(self, ctx):
        self.ctx = ctx  # AppContext: listings, products, accounts, compatibility, catalog

    # ------------------------------------------------------------------ prepared listings
    def export_listings(self, listings: list[Listing], target: Path, copy_photos: bool = True) -> ExportResult:
        result = ExportResult(target)
        photos_dir = target.with_name(target.stem + "-fotos") if copy_photos else None
        if photos_dir is not None:
            photos_dir.mkdir(parents=True, exist_ok=True)
            result.photos_dir = photos_dir
        rows: list[list] = []
        issues: list[list] = []
        copied: dict[Path, str] = {}
        for listing in listings:
            product = self.ctx.products.get(listing.product_id)
            if product is None:
                continue
            payload = self.ctx.listings.payload(listing, product)
            report = self.ctx.listings.validate(listing, product)
            account = self.ctx.accounts.get(listing.account_id)
            status = listing.status.label
            result.listings += 1
            result.by_status[status] += 1
            errors = [i for i in report.issues if i.severity == "error"]
            if errors:
                result.with_errors += 1
            for issue in report.issues:
                issues.append([product.sku, payload.title or product.internal_name, status,
                               "Bloqueia" if issue.severity == "error" else "Aviso", issue.field, issue.message])
            pending = "; ".join(i.message for i in report.issues)
            photo_names = []
            for n, source in enumerate(payload.images, start=1):
                if photos_dir is None:
                    photo_names.append(source.name)
                    continue
                if source in copied:
                    photo_names.append(copied[source])
                    continue
                if not source.is_file():
                    result.images_missing += 1
                    continue
                name = f"{_file_safe(product.sku)}-{n:02d}{source.suffix.lower()}"
                shutil.copy2(source, photos_dir / name)
                copied[source] = name
                result.images_copied += 1
                photo_names.append(name)
            compat = self.ctx.compatibility.list_for_product(product.id)
            confirmed = [_compat_text(c) for c in compat if c.verification is Verification.CONFIRMED]
            unverified = sum(1 for c in compat if c.verification is Verification.UNVERIFIED)
            attributes = "; ".join(f"{k}: {v}" for k, v in payload.attributes.items())
            base = {
                "status": status, "pending": pending, "marketplace": listing.marketplace.label,
                "account": account.label if account else "", "title": payload.title,
                "description": payload.description, "condition": product.condition.label, "gtin": product.gtin,
                "brand": product.brand, "model": product.model, "category": product.category,
                "subcategory": product.subcategory, "category_id": payload.category_id,
                "listing_type": payload.listing_type, "weight": product.weight_g or "",
                "length": _cm(product.length_mm), "width": _cm(product.width_mm), "height": _cm(product.height_mm),
                "photos": "; ".join(photo_names), "attributes": attributes, "compat": "; ".join(confirmed),
                "unverified": unverified or "", "source": product.source or "Cadastro manual",
            }
            if product.variants:
                for variant in product.variants:
                    price = variant.price_cents if variant.price_cents is not None else payload.price_cents
                    rows.append(_listing_row(base, variant.sku or "", product.sku, variant.name,
                                             "; ".join(f"{k}: {v}" for k, v in variant.attributes.items()),
                                             price, variant.stock))
            else:
                rows.append(_listing_row(base, product.sku, "", "", "", payload.price_cents, payload.stock))
        result.rows = len(rows)
        writer = SpreadsheetWriter(target)
        writer.add("Anúncios", LISTING_COLUMNS, rows)
        if target.suffix.lower() != ".csv":
            writer.add("Pendências", ISSUE_COLUMNS, issues)
            summary = [["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")],
                       ["Anúncios exportados", result.listings], ["Linhas (uma por variação)", result.rows],
                       ["Anúncios com pendência bloqueante", result.with_errors]]
            summary += [[f"Situação: {name}", n] for name, n in sorted(result.by_status.items())]
            if photos_dir is not None:
                summary += [["Pasta das fotos", str(photos_dir)], ["Fotos copiadas", result.images_copied],
                            ["Fotos não encontradas", result.images_missing]]
            summary += [["Observação", NOTICE]]
            writer.add("Resumo", ["Item", "Valor"], summary)
        writer.save()
        log.info("exportação de anúncios: %s (%s)", target, result.summary())
        return result

    # ------------------------------------------------------------------ catalog pending report
    def export_catalog_issues(self, target: Path) -> ExportResult:
        """Every product with its validation issues and pending media (what the client still has to send)."""
        result = ExportResult(target)
        rows: list[list] = []
        for summary in self.ctx.catalog.search():
            product = self.ctx.products.get(summary.id)
            if product is None:
                continue
            result.listings += 1
            result.by_status[product.status.label] += 1
            report = self.ctx.catalog.validate(product)
            problems = [(("Bloqueia" if i.severity == "error" else "Aviso"), i.field, i.message)
                        for i in report.issues]
            problems += [("Aviso", "mídia", f"Mídia citada na importação não encontrada: {m.reference}")
                         for m in product.missing_media]
            if any(level == "Bloqueia" for level, _, _ in problems):
                result.with_errors += 1
            for level, field_name, message in problems:
                rows.append([product.sku, product.internal_name, product.status.label, level, field_name, message])
        result.rows = len(rows)
        writer = SpreadsheetWriter(target)
        writer.add("Pendências", ISSUE_COLUMNS, rows)
        if target.suffix.lower() != ".csv":
            summary = [["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")],
                       ["Produtos no catálogo", result.listings],
                       ["Produtos com pendência bloqueante", result.with_errors]]
            summary += [[f"Situação: {name}", n] for name, n in sorted(result.by_status.items())]
            writer.add("Resumo", ["Item", "Valor"], summary)
        writer.save()
        log.info("relatório de pendências do catálogo: %s (%s linhas)", target, result.rows)
        return result


def _listing_row(base: dict, sku: str, parent_sku: str, variant: str, variant_attrs: str, price: int | None,
                 stock: int | None) -> list:
    return [base["status"], base["pending"], base["marketplace"], base["account"], sku, parent_sku, variant,
            variant_attrs, base["title"], base["description"], format_money(price, symbol=False) if price else "",
            stock if stock is not None else "", base["condition"], base["gtin"], base["brand"], base["model"],
            base["category"], base["subcategory"], base["category_id"], base["listing_type"], base["weight"],
            base["length"], base["width"], base["height"], base["photos"], base["attributes"], base["compat"],
            base["unverified"], base["source"]]


def _compat_text(c) -> str:
    parts = [p for p in (c.brand, c.model, c.version) if p]
    years = c.years()
    return " ".join(parts) + (f" ({years})" if years else "")


def _file_safe(text: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in text)[:60] or "produto"


__all__ = ["ExportResult", "ListingExporter", "NOTICE", "SpreadsheetWriter"]
