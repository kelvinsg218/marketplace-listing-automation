"""Export actions on the Anúncios and Catálogo pages."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QFileDialog

from mla.core.enums import Marketplace
from mla.ui.main_window import MainWindow
from mla.ui.theme import apply_theme


@pytest.fixture
def window(qtbot, ctx):
    apply_theme(QApplication.instance())
    win = MainWindow(ctx, start_executor=False)
    qtbot.addWidget(win)
    win.show()
    yield win
    win.services.executor.stop(timeout=1)


def test_listings_page_exports_current_filter(window, ctx, complete_product, tmp_path, monkeypatch):
    from mla.ui.pages import listings as listings_page
    product = complete_product("UX-1")
    ctx.catalog.save(product)
    account = ctx.accounts.default_for(Marketplace.SIMULATED)
    ctx.listings.prepare([product.id], Marketplace.SIMULATED, account.id)
    target = tmp_path / "anuncios.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(target), "")))
    shown = []
    monkeypatch.setattr(listings_page, "inform", lambda _p, title, text, error=False: shown.append((title, text)))
    window.navigate("listings")
    page = window.pages["listings"]
    page.marketplace.setCurrentIndex(page.marketplace.findData(Marketplace.SIMULATED))
    page._export()
    assert shown and shown[0][0] == "Planilha exportada" and "1 anúncio(s) em 1 linha(s)" in shown[0][1]
    rows = list(load_workbook(target)["Anúncios"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and "UX-1" in rows[0]


def test_catalog_issues_report_from_products_menu(window, ctx, complete_product, tmp_path, monkeypatch):
    ctx.catalog.save(complete_product("UX-2", description=""))
    target = tmp_path / "pendencias.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(target), "")))
    window.navigate("products")
    window.pages["products"]._export_issues()
    rows = list(load_workbook(target)["Pendências"].iter_rows(min_row=2, values_only=True))
    assert any(r[0] == "UX-2" for r in rows)
