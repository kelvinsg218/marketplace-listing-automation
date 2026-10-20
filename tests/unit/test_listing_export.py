"""Intermediate listings spreadsheet and catalog pending-issues report."""
import csv

from openpyxl import load_workbook

from mla.core.enums import DataSource, Marketplace, Verification
from mla.core.models import Compatibility, ProductVariant
from mla.services.listing_export import LISTING_COLUMNS, ListingExporter

SIM = Marketplace.SIMULATED


def _prepared(ctx, product):
    ctx.catalog.save(product)
    account = ctx.accounts.default_for(SIM)
    return ctx.listings.get(ctx.listings.prepare([product.id], SIM, account.id).created[0])


def _column(row, name):
    return row[LISTING_COLUMNS.index(name)]


def test_xlsx_has_one_row_per_variation_photos_and_only_confirmed_compat(ctx, complete_product, tmp_path):
    product = complete_product("EX-1", gtin="", title="=HYPERLINK(\"http://x\")")
    product.variants = [ProductVariant(name="Preto", sku="EX-1-P", stock=3, attributes={"Cor": "Preto"}),
                        ProductVariant(name="Branco", sku="EX-1-B", stock=0, price_cents=13990)]
    listing = _prepared(ctx, product)
    ctx.compatibility.save(Compatibility(product.id, "Onix", brand="Chevrolet", year_from=2019, year_to=2022,
                                         verification=Verification.CONFIRMED))
    ctx.compatibility.save(Compatibility(product.id, "Gol", brand="VW", source=DataSource.EXTERNAL))
    target = tmp_path / "saida" / "anuncios.xlsx"
    result = ListingExporter(ctx).export_listings([listing], target)
    assert (result.listings, result.rows, result.images_copied) == (1, 2, 1)
    book = load_workbook(target)
    assert book.sheetnames == ["Anúncios", "Pendências", "Resumo"]
    rows = list(book["Anúncios"].iter_rows(values_only=True))
    assert list(rows[0]) == LISTING_COLUMNS
    black, white = rows[1], rows[2]
    assert (_column(black, "SKU"), _column(black, "SKU do produto pai"), _column(black, "Variação")) == \
        ("EX-1-P", "EX-1", "Preto")
    assert _column(black, "Preço (R$)") == "129,90" and _column(white, "Preço (R$)") == "139,90"
    assert _column(black, "Estoque") == 3 and _column(white, "Estoque") == 0
    assert _column(black, "GTIN/EAN") in (None, "")  # never invented
    assert _column(black, "Compatibilidades confirmadas") == "Chevrolet Onix (2019–2022)"
    assert _column(black, "Compatibilidades não verificadas (não exportadas)") == 1
    assert _column(black, "Título") == "=HYPERLINK(\"http://x\")"  # stored as text, not as a formula
    assert book["Anúncios"].cell(row=2, column=LISTING_COLUMNS.index("Título") + 1).data_type == "s"
    photo = _column(black, "Fotos (arquivos)")
    assert photo == "EX-1-01.png" and (tmp_path / "saida" / "anuncios-fotos" / photo).is_file()
    summary = {r[0]: r[1] for r in book["Resumo"].iter_rows(min_row=2, values_only=True)}
    assert summary["Anúncios exportados"] == 1 and "não é o modelo oficial" in summary["Observação"].lower()


def test_csv_export_neutralises_formulas(ctx, complete_product, tmp_path):
    listing = _prepared(ctx, complete_product("EX-2", title="+CMD|' /C calc'!A0"))
    target = tmp_path / "anuncios.csv"
    result = ListingExporter(ctx).export_listings([listing], target, copy_photos=False)
    assert result.rows == 1 and result.photos_dir is None
    with target.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    assert rows[1][LISTING_COLUMNS.index("Título")].startswith("'+CMD")


def test_pending_listing_is_exported_with_its_problems(ctx, complete_product, tmp_path):
    listing = _prepared(ctx, complete_product("EX-3", description=""))
    result = ListingExporter(ctx).export_listings([listing], tmp_path / "a.xlsx", copy_photos=False)
    assert result.with_errors == 1
    book = load_workbook(tmp_path / "a.xlsx")
    issues = list(book["Pendências"].iter_rows(min_row=2, values_only=True))
    assert any(r[0] == "EX-3" and r[3] == "Bloqueia" and "descrição" in r[5] for r in issues)


def test_catalog_issues_report_lists_blocking_problems_and_missing_media(ctx, complete_product, tmp_path):
    from mla.core.models import MissingMedia, Product
    ctx.catalog.save(complete_product("OK-1"))
    incomplete = Product(internal_name="Sem preço", sku="INC-1")
    incomplete.missing_media = [MissingMedia("foto-inc-1.jpg")]
    ctx.catalog.save(incomplete)
    result = ListingExporter(ctx).export_catalog_issues(tmp_path / "pendencias.xlsx")
    assert result.listings == 2 and result.with_errors == 1
    rows = list(load_workbook(tmp_path / "pendencias.xlsx")["Pendências"].iter_rows(min_row=2, values_only=True))
    assert all(r[0] != "OK-1" or r[3] == "Aviso" for r in rows)
    assert any(r[0] == "INC-1" and "foto-inc-1.jpg" in r[5] for r in rows)
